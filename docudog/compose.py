"""Build a third document from already-classified sources (brief or handover)."""

from __future__ import annotations

import logging
import os
import re
from datetime import datetime, timezone
from typing import Any

from . import artifact_home, cadence, file_ids, inference, reporter, router
from .paths_util import normalize_fs_path
from .security_labels import format_security_level

logger = logging.getLogger(__name__)

MAX_SOURCES_DEFAULT = 8
MAX_CHARS_PER_SOURCE_DEFAULT = 4000
KIND_BRIEF = "brief"
KIND_HANDOVER = "handover"
_KINDS = frozenset({KIND_BRIEF, KIND_HANDOVER})
_DRAFT_MARKERS = (
    "_초안",
    "-초안",
    "_draft",
    "-draft",
    "_wip",
    "_임시",
    "_임시본",
)
_FINAL_MARKERS = (
    "_최종",
    "-최종",
    "_final",
    "-final",
    "_완료",
    "_확정",
)
_SENSITIVE = frozenset({"P1", "P2"})


def compose_settings(cfg: dict[str, Any] | None) -> dict[str, Any]:
    raw = (cfg or {}).get("compose_settings")
    block = raw if isinstance(raw, dict) else {}
    try:
        max_sources = int(block.get("max_sources", MAX_SOURCES_DEFAULT))
    except (TypeError, ValueError):
        max_sources = MAX_SOURCES_DEFAULT
    try:
        max_chars = int(block.get("max_chars_per_source", MAX_CHARS_PER_SOURCE_DEFAULT))
    except (TypeError, ValueError):
        max_chars = MAX_CHARS_PER_SOURCE_DEFAULT
    formats = block.get("formats_enabled")
    enabled_formats = ["md", "html", "docx"]
    if isinstance(formats, list) and formats:
        enabled_formats = [str(x).strip().lower() for x in formats if str(x).strip()]
    return {
        "max_sources": max(1, min(20, max_sources)),
        "max_chars_per_source": max(200, min(20000, max_chars)),
        "allow_p1_body": bool(block.get("allow_p1_body", False)),
        "formats_enabled": enabled_formats,
    }


def composed_dir(cfg: dict[str, Any] | None) -> str:
    paths = (cfg or {}).get("paths") if isinstance((cfg or {}).get("paths"), dict) else {}
    raw = str((paths or {}).get("composed_dir") or "").strip()
    if raw:
        return normalize_fs_path(os.path.expandvars(raw))
    return os.path.join(artifact_home.artifact_home(cfg), "composed")


def body_allowed(security_level: str, cfg: dict[str, Any] | None) -> bool:
    code = str(security_level or "").strip().upper()
    if code not in _SENSITIVE:
        return True
    return bool(compose_settings(cfg).get("allow_p1_body"))


def is_draft_name(path: str) -> bool:
    stem = os.path.splitext(os.path.basename(path))[0].casefold()
    return any(m.casefold() in stem for m in _DRAFT_MARKERS)


def is_final_name(path: str) -> bool:
    stem = os.path.splitext(os.path.basename(path))[0].casefold()
    return any(m.casefold() in stem for m in _FINAL_MARKERS)


def _file_mtime_iso(path: str) -> str:
    try:
        ts = os.path.getmtime(path)
    except OSError:
        return ""
    return datetime.fromtimestamp(ts, tz=timezone.utc).isoformat()


def _stale(path: str, meta: dict[str, Any]) -> bool:
    analyzed = str(meta.get("last_analyzed_utc") or "").strip()
    mtime = _file_mtime_iso(path)
    if not analyzed or not mtime:
        return False
    return mtime[:19] > analyzed[:19]


def iter_classified(state: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    files = state.get("files") if isinstance(state.get("files"), dict) else {}
    rows: list[tuple[str, dict[str, Any]]] = []
    for path, meta in files.items():
        if not isinstance(meta, dict):
            continue
        if not str(meta.get("summary") or meta.get("security_level") or "").strip():
            continue
        file_ids.ensure_file_id(meta)
        rows.append((str(path), meta))
    rows.sort(
        key=lambda pm: str(pm[1].get("last_analyzed_utc") or ""),
        reverse=True,
    )
    return rows


def candidate_rows(cfg: dict[str, Any], state: dict[str, Any], *, limit: int = 40) -> list[dict[str, Any]]:
    """Recent classified files for the operator checklist."""
    out: list[dict[str, Any]] = []
    seen_sha: set[str] = set()
    for path, meta in iter_classified(state):
        sha = str(meta.get("sha256") or "")
        if sha and sha in seen_sha:
            continue
        if sha:
            seen_sha.add(sha)
        sec = str(meta.get("security_level") or "").upper()
        out.append(
            {
                "file_id": str(meta.get("file_id") or ""),
                "path": path,
                "basename": os.path.basename(path),
                "security_level": sec,
                "security_label": format_security_level(sec, cfg),
                "summary": str(meta.get("summary") or "").strip(),
                "draft": is_draft_name(path),
                "final": is_final_name(path),
                "stale": _stale(path, meta),
                "exists": os.path.isfile(path),
                "body_ok": body_allowed(sec, cfg),
            }
        )
        if len(out) >= limit:
            break
    return out


def thread_options(state: dict[str, Any]) -> list[dict[str, str]]:
    rows = state.get("threads") if isinstance(state.get("threads"), list) else []
    opts: list[dict[str, str]] = []
    for th in rows[:12]:
        if not isinstance(th, dict):
            continue
        tid = str(th.get("id") or "").strip()
        title = str(th.get("title") or tid).strip()
        if tid:
            opts.append({"id": tid, "title": title})
    return opts


def file_ids_for_thread(state: dict[str, Any], thread_id: str) -> list[str]:
    want = str(thread_id or "").strip()
    if not want:
        return []
    rows = state.get("threads") if isinstance(state.get("threads"), list) else []
    files = state.get("files") if isinstance(state.get("files"), dict) else {}
    for th in rows:
        if not isinstance(th, dict) or str(th.get("id") or "") != want:
            continue
        ids: list[str] = []
        for mem in th.get("members") or []:
            if not isinstance(mem, dict):
                continue
            path = str(mem.get("path") or "")
            meta = files.get(path)
            if isinstance(meta, dict):
                ids.append(file_ids.ensure_file_id(meta))
            elif mem.get("file_id"):
                ids.append(str(mem.get("file_id")))
        return ids
    return []


def resolve_sources(
    cfg: dict[str, Any],
    state: dict[str, Any],
    *,
    selected_ids: list[str],
    selected_paths: list[str] | None = None,
) -> tuple[list[dict[str, Any]], str | None]:
    settings = compose_settings(cfg)
    files = state.get("files") if isinstance(state.get("files"), dict) else {}
    picked: list[dict[str, Any]] = []
    seen: set[str] = set()

    def _add(path: str, meta: dict[str, Any]) -> None:
        key = os.path.normcase(path)
        if key in seen:
            return
        seen.add(key)
        sec = str(meta.get("security_level") or "").upper()
        skip_reason = str(meta.get("skip_reason") or "")
        picked.append(
            {
                "file_id": file_ids.ensure_file_id(meta),
                "path": path,
                "basename": os.path.basename(path),
                "ext": os.path.splitext(path)[1].lower(),
                "security_level": sec,
                "security_label": format_security_level(sec, cfg),
                "tags": meta.get("tags") if isinstance(meta.get("tags"), list) else [],
                "summary": str(meta.get("summary") or "").strip(),
                "last_analyzed_utc": str(meta.get("last_analyzed_utc") or ""),
                "file_mtime_utc": _file_mtime_iso(path),
                "stale": _stale(path, meta),
                "sha256": str(meta.get("sha256") or ""),
                "draft": is_draft_name(path),
                "final": is_final_name(path),
                "body_ok": body_allowed(sec, cfg),
                "exists": os.path.isfile(path),
                "skip_extract": skip_reason,
            }
        )

    for fid in selected_ids:
        path = file_ids.find_path_by_file_id(files, fid)
        if not path:
            continue
        meta = files.get(path)
        if isinstance(meta, dict):
            _add(path, meta)
    for raw in selected_paths or []:
        path = normalize_fs_path(os.path.expandvars(str(raw)))
        meta = files.get(path)
        if not isinstance(meta, dict):
            for p, m in files.items():
                if isinstance(m, dict) and os.path.normcase(p) == os.path.normcase(path):
                    meta = m
                    path = p
                    break
        if isinstance(meta, dict):
            _add(path, meta)

    if not picked:
        return [], "분류된 파일을 하나도 고르지 못했음. 이미 분석된 항목만 만들기에 쓸 수 있음."
    if len(picked) > settings["max_sources"]:
        return [], f"너무 많음. {settings['max_sources']}개까지 고르세요."
    return picked, None


def _excerpt(path: str, cfg: dict[str, Any], cap: int) -> tuple[str, str]:
    text, skip = router.extract_document_text(path, cfg)
    if skip:
        return "", skip
    if not text:
        return "", "empty"
    clipped = text.strip()[:cap]
    return clipped, ""


def _source_notes(cfg: dict[str, Any], sources: list[dict[str, Any]]) -> list[str]:
    settings = compose_settings(cfg)
    cap = int(settings["max_chars_per_source"])
    notes: list[str] = []
    for src in sources:
        head = (
            f"{src['basename']} | {src['security_label'] or src['security_level']} | "
            f"{'이전 버전' if src['draft'] and not src['final'] else '후보'}"
        )
        if src["stale"]:
            head += " | 분류가 파일보다 오래됨"
        if not src["exists"]:
            notes.append(head + "\n(파일이 없음)")
            continue
        if not src["body_ok"]:
            notes.append(
                head
                + f"\n요약: {src['summary'] or '(없음)'}\n본문: 민감 등급이라 내용 없이 목록만."
            )
            continue
        body, skip = _excerpt(src["path"], cfg, cap)
        if skip or not body:
            notes.append(
                head + f"\n요약: {src['summary'] or '(없음)'}\n본문 없음: {skip or 'empty'}"
            )
            continue
        notes.append(head + f"\n요약: {src['summary'] or '(없음)'}\n발췌:\n{body}")
    return notes


def _llm_markdown(cfg: dict[str, Any], system: str, user: str) -> str:
    raw = inference.run_aux_completion(
        system,
        user,
        cfg,
        max_tokens=700,
        task_log_label="compose",
    )
    text = (raw or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:markdown|md)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _heuristic_brief(cfg: dict[str, Any], intent: str, sources: list[dict[str, Any]]) -> str:
    settings = compose_settings(cfg)
    cap = min(400, int(settings["max_chars_per_source"]))
    lines = [f"요청: {intent.strip() or '요약'}", ""]
    for src in sources:
        bit = src["summary"] or "(요약 없음)"
        extra = []
        if src["draft"] and not src["final"]:
            extra.append("이전 버전")
        if src["final"]:
            extra.append("최종 파일명")
        if not src["body_ok"]:
            extra.append("민감 — 본문 없음")
        suffix = f" ({', '.join(extra)})" if extra else ""
        lines.append(f"- **{src['basename']}**{suffix}: {bit}")
        if src["body_ok"] and src["exists"]:
            body, skip = _excerpt(src["path"], cfg, cap)
            if body:
                lines.append(f"  - 발췌: {body.replace(chr(10), ' ')[:400]}")
            elif skip:
                lines.append(f"  - 본문 없음 ({skip})")
    lines.append("")
    lines.append("연습 모드이거나 모델이 비어 있어, 분류 요약과 짧은 발췌만 이어 붙였음.")
    return "\n".join(lines)


def _sort_handover(sources: list[dict[str, Any]]) -> list[dict[str, Any]]:
    def key(src: dict[str, Any]) -> tuple[int, int, str]:
        if src["final"]:
            bucket = 0
        elif src["draft"]:
            bucket = 2
        else:
            bucket = 1
        return (bucket, 0 if src["body_ok"] else 1, src["last_analyzed_utc"] or "")

    return sorted(sources, key=key)


def _heuristic_handover(
    cfg: dict[str, Any],
    sources: list[dict[str, Any]],
    state: dict[str, Any],
) -> str:
    ordered = _sort_handover(sources)
    latest = [s for s in ordered if s["final"] or not s["draft"]][:2]
    if not latest:
        latest = ordered[:2]
    tags: list[str] = []
    for s in sources:
        for t in s.get("tags") or []:
            ts = str(t).strip()
            if ts and ts not in tags:
                tags.append(ts)
    overview_bits = [s["summary"] for s in sources if s["summary"]][:3]
    overview = " / ".join(overview_bits) if overview_bits else "고른 파일들의 업무 묶음."
    if tags:
        overview += " 태그: " + ", ".join(tags[:8])

    lines = [
        "## 개요",
        "",
        overview,
        "",
        "지금 볼 파일:",
    ]
    for s in latest:
        lines.append(f"- `{s['path']}` — {s['summary'] or s['basename']}")
    if any(not s["body_ok"] for s in sources):
        lines.append("")
        lines.append("민감 문서는 목록만 있고 내용은 합성하지 않음.")
    lines.extend(["", "## 읽는 순서", ""])
    why = {
        0: "최신·확정으로 보이는 산출물부터.",
        1: "맥락·참고.",
        2: "이전 버전. 최신을 본 뒤에만.",
    }
    for i, s in enumerate(ordered, start=1):
        bucket = 0 if s["final"] else (2 if s["draft"] else 1)
        lines.append(f"{i}. `{s['basename']}` — {why[bucket]}")
    lines.extend(["", "## 문서 목록", ""])
    lines.append("| 이름 | 경로 | 형식 | 등급 | 분류 시각 | 상태 | 한 줄 |")
    lines.append("|---|---|---|---|---|---|---|")
    for s in ordered:
        status = []
        if s["draft"] and not s["final"]:
            status.append("이전 버전")
        if s["final"]:
            status.append("최종 파일명")
        if s["stale"]:
            status.append("다시 분류 대기")
        if not s["exists"]:
            status.append("파일 없음")
        if not s["body_ok"]:
            status.append("본문 생략")
        lines.append(
            "| {name} | `{path}` | {ext} | {sec} | {an} | {st} | {sum} |".format(
                name=s["basename"].replace("|", " "),
                path=s["path"].replace("|", " "),
                ext=s["ext"] or "-",
                sec=s["security_label"] or s["security_level"] or "-",
                an=(s["last_analyzed_utc"] or "-")[:19],
                st=", ".join(status) or "-",
                sum=(s["summary"] or "-").replace("|", " ")[:120],
            )
        )
    gaps: list[str] = []
    for s in sources:
        if not s["exists"]:
            gaps.append(f"- `{s['basename']}` 경로에 파일이 없음.")
        elif s.get("skip_extract"):
            gaps.append(f"- `{s['basename']}` 본문을 못 읽음.")
        elif not s["body_ok"]:
            gaps.append(f"- `{s['basename']}` 민감 등급이라 내용 합성 안 함.")
    for row in cadence.evaluate_cadence(cfg, state):
        if row.get("status") == "miss":
            gaps.append(f"- 주기 문서 공백: {row.get('message')}")
    lines.extend(["", "## 공백", ""])
    if gaps:
        lines.extend(gaps)
    else:
        lines.append("- 이 묶음에서 표시할 공백 없음.")
    return "\n".join(lines) + "\n"


def _wrap_document(
    cfg: dict[str, Any],
    *,
    kind: str,
    intent: str,
    sources: list[dict[str, Any]],
    body: str,
    backend: str,
) -> str:
    now = reporter._format_report_timestamp(datetime.now(timezone.utc))
    title = "인수인계 / OJT" if kind == KIND_HANDOVER else "요약·보고서"
    src_lines = []
    for s in sources:
        flag = []
        if s["stale"]:
            flag.append("stale")
        if not s["body_ok"]:
            flag.append("본문생략")
        extra = f" ({', '.join(flag)})" if flag else ""
        src_lines.append(f"- `{s['path']}`{extra}")
    return (
        f"# DocuDog {title}\n\n"
        f"- 생성: {now}\n"
        f"- 모델: `{backend or 'unknown'}`\n"
        f"- 요청: {intent.strip() or '(없음)'}\n\n"
        "## 사용한 원본\n\n"
        + "\n".join(src_lines)
        + "\n\n"
        "> 이 파일은 원본을 대체하지 않음. 보안 등급은 참고용임. "
        "민감(P1/P2) 본문은 기본으로 넣지 않음.\n\n"
        + body.strip()
        + "\n"
    )


def _write_docx(md_path: str, docx_path: str) -> str | None:
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx missing; skip Word output")
        return None
    doc = Document()
    try:
        with open(md_path, encoding="utf-8") as f:
            lines = f.read().splitlines()
        in_table = False
        for line in lines:
            s = line.strip()
            if s.startswith("|"):
                if not in_table:
                    in_table = True
                cells = [c.strip().strip("`") for c in s.strip("|").split("|")]
                if cells and set(cells[0]) <= set("-:"):
                    continue
                doc.add_paragraph(" | ".join(cells))
                continue
            in_table = False
            if s.startswith("# "):
                doc.add_heading(s[2:].strip(), level=0)
            elif s.startswith("## "):
                doc.add_heading(s[3:].strip(), level=1)
            elif s.startswith("### "):
                doc.add_heading(s[4:].strip(), level=2)
            elif s.startswith("> "):
                doc.add_paragraph(s[2:].strip())
            elif s.startswith("- "):
                doc.add_paragraph(s[2:].strip(), style="List Bullet")
            elif s:
                doc.add_paragraph(s)
        os.makedirs(os.path.dirname(docx_path) or ".", exist_ok=True)
        doc.save(docx_path)
    except OSError as e:
        logger.warning("DOCX write failed: %s", e)
        return None
    return docx_path


def _slug(kind: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    return f"{stamp}-{kind}"


def compose_markdown(
    cfg: dict[str, Any],
    state: dict[str, Any],
    *,
    kind: str,
    intent: str,
    sources: list[dict[str, Any]],
) -> str:
    kind_n = kind if kind in _KINDS else KIND_BRIEF
    backend = str(state.get("last_inference_backend") or "")
    model_cfg = cfg.get("model") if isinstance(cfg.get("model"), dict) else {}
    if bool(model_cfg.get("use_mock", True)):
        backend = backend or "mock"
    notes = _source_notes(cfg, sources)
    joined = "\n\n---\n\n".join(notes)
    if kind_n == KIND_HANDOVER:
        heuristic = _heuristic_handover(cfg, sources, state)
        system = (
            "You write a Korean handover/OJT markdown with exactly these headings: "
            "## 개요, ## 읽는 순서, ## 문서 목록, ## 공백. "
            "Do not invent file paths. Keep P1/P2 as list-only. "
            "Put draft/_초안 files after finals. Output markdown only."
        )
        user = f"요청: {intent}\n\n고정 뼈대(이걸 다듬되 절은 유지):\n{heuristic}\n\n소스 메모:\n{joined}"
        llm = _llm_markdown(cfg, system, user)
        body = llm if llm else heuristic
    else:
        heuristic = _heuristic_brief(cfg, intent, sources)
        system = (
            "You write a short Korean briefing from the source notes. "
            "Honor the user request. Do not include P1/P2 body text. "
            "Do not claim to replace the originals. Markdown only, no title heading."
        )
        user = f"요청: {intent}\n\n소스 메모:\n{joined}"
        llm = _llm_markdown(cfg, system, user)
        body = llm if llm else heuristic
    return _wrap_document(
        cfg, kind=kind_n, intent=intent, sources=sources, body=body, backend=backend or "mock"
    )


def write_outputs(
    cfg: dict[str, Any],
    markdown: str,
    *,
    kind: str,
    formats: list[str] | None = None,
) -> dict[str, str]:
    settings = compose_settings(cfg)
    allowed = set(settings["formats_enabled"]) | {"md"}
    want = [f.strip().lower() for f in (formats or ["md", "html"]) if f.strip()]
    if not want:
        want = ["md"]
    out_dir = composed_dir(cfg)
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.join(out_dir, _slug(kind))
    md_path = base + ".md"
    with open(md_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(markdown)
    written: dict[str, str] = {"md": md_path}
    if "html" in want and "html" in allowed:
        html_path = reporter.sync_report_html(md_path)
        if html_path:
            written["html"] = html_path
    if "docx" in want and "docx" in allowed:
        docx_path = _write_docx(md_path, base + ".docx")
        if docx_path:
            written["docx"] = docx_path
    return written


def compose(
    cfg: dict[str, Any],
    state: dict[str, Any],
    *,
    kind: str = KIND_BRIEF,
    intent: str = "",
    file_ids_selected: list[str] | None = None,
    paths_selected: list[str] | None = None,
    thread_id: str = "",
    formats: list[str] | None = None,
    report_path: str = "",
) -> dict[str, Any]:
    """
    Create a third document. Returns {ok, error?, paths?, markdown?}.
    Does not move or rename originals.
    """
    ids = list(file_ids_selected or [])
    if thread_id and not ids:
        ids = file_ids_for_thread(state, thread_id)
    sources, err = resolve_sources(
        cfg, state, selected_ids=ids, selected_paths=paths_selected
    )
    if err:
        return {"ok": False, "error": err}
    kind_n = kind if kind in _KINDS else KIND_BRIEF
    markdown = compose_markdown(cfg, state, kind=kind_n, intent=intent, sources=sources)
    paths = write_outputs(cfg, markdown, kind=kind_n, formats=formats)
    rp = report_path or artifact_home.resolve_report_path(cfg)
    try:
        from . import activity as activity_mod

        names = ", ".join(s["basename"] for s in sources)
        activity_mod.append_activity(
            cfg,
            rp,
            "compose",
            f"{kind_n} sources={len(sources)} [{names}] -> {paths.get('md', '')}",
        )
    except Exception:
        logger.debug("compose activity log failed", exc_info=True)
    return {"ok": True, "paths": paths, "markdown": markdown, "sources": sources}
